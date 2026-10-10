#!/usr/bin/env python3
"""把設定檔 xlsx 重新打包成較小的檔案（內容不變）。

1. 只保留儲存格值、標題列樣式、欄寬、凍結窗格、分頁顯示狀態；去掉其他格式殘留。
2. 把 XML 中的 &#NNNN; 字元實體改成 UTF-8 原字（中文每字由約 8 位元組降到 3 位元組）。
3. 以最高壓縮率重新壓縮。
不保留的後設資訊（2026-10-10 對 PR #76 前的正式檔逐項核對）：
- 命名範圍：原檔 295 個，全部是 Excel 自動篩選留下的隱藏 `_xlnm._FilterDatabase`，只記錄篩選範圍，無資料語義，故不保留。
- 資料驗證、條件格式、儲存格註解、超連結、表格物件：原檔皆為 0 個。
- 其他儲存格格式（字型、顏色、對齊）與篩選狀態：不保留。
因此「內容不變」指每個儲存格的值與分頁顯示狀態不變（由 config_export.py --check 的 CSV 逐格比對保證），
不是指整個 xlsx 位元組或格式完全相同。若日後設定檔加入資料驗證或註解，執行本工具前須先擴充保留邏輯。
用法：python3 tools/config_compact.py [檔案路徑]（預設 config/世界模擬遊戲設定檔.xlsx）
"""
import io, pathlib, re, sys, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
path = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "config" / "世界模擬遊戲設定檔.xlsx"

src = openpyxl.load_workbook(path)
dst = openpyxl.Workbook()
dst.remove(dst.active)
hdr_font, hdr_fill = Font(bold=True), PatternFill("solid", fgColor="DDE6F0")
for ws in src.worksheets:
    n = dst.create_sheet(ws.title)
    for row in ws.iter_rows(values_only=True):
        n.append(list(row))
    while n.max_row > 1 and all(c.value in (None, "") for c in n[n.max_row]):
        n.delete_rows(n.max_row)
    for c in n[1]:
        c.font, c.fill = hdr_font, hdr_fill
    for k, cd in ws.column_dimensions.items():
        if cd.width:
            n.column_dimensions[k].width = cd.width
    n.freeze_panes = ws.freeze_panes or "A2"
    n.sheet_state = ws.sheet_state
dst.active = 0
buf = io.BytesIO()
dst.save(buf)

ent = re.compile(rb"&#(\d+);")
def unescape(m):
    cp = int(m.group(1))
    return m.group(0) if cp < 128 else chr(cp).encode("utf-8")

out = io.BytesIO()
with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
            data = ent.sub(unescape, data)
        zout.writestr(item.filename, data)
path.write_bytes(out.getvalue())
print(f"{path.name}: {len(out.getvalue()):,} 位元組")
